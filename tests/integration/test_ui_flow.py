"""The board as the browser drives it: the UI's calls in the UI's order, through the Vite proxy without a token
(the proxy adds it), ending with a reload that must bring every piece of drawer state back.

test_stack.py checks each route; this checks the sequence the UI depends on (frontend/src/stores/cases.ts).
"""

import io
import time

import httpx
import pytest
from docx import Document
from pptx import Presentation

from navigator.settings import PROJECT_ROOT

pytestmark = pytest.mark.integration

GLUCORA_PDF = PROJECT_ROOT / "docs" / "examples" / "glucora-business-model.pdf"
POLL_S = 1.5  # the store's poll interval
TIMEOUT_S = 600
NOTE = "Internal: call them before Friday."  # notes are internal: never in a pitch or a document


@pytest.fixture
def ui(stack) -> httpx.Client:
    """The browser's view: the frontend's origin, no Authorization header."""
    return httpx.Client(base_url=f"{stack.ui}/api", timeout=180)


def poll(ui: httpx.Client, path: str, key: str) -> dict:
    body = ui.get(path).json()
    deadline = time.monotonic() + TIMEOUT_S
    while body[key] == "pending":
        assert time.monotonic() < deadline, f"{path} still pending"
        time.sleep(POLL_S)
        body = ui.get(path).json()
    return body


def test_board_flow_through_the_proxy(ui):
    # Board: earlier Glucora cases (from other tests) should come back as similar cases
    earlier = [c["id"] for c in ui.get("/cases").json() if c["company_name"] == "Glucora Health"]
    case = ui.post("/cases", files={"file": (GLUCORA_PDF.name, GLUCORA_PDF.read_bytes(), "application/pdf")}).json()
    cid = case["id"]
    case = poll(ui, f"/cases/{cid}", "analysis_status")
    assert case["analysis_status"] == "done" and case["company_name"] == "Glucora Health"
    if earlier:
        assert {s["case_id"] for s in case["similar_cases"]} & set(earlier)

    # Open the drawer: the store loads these together
    refs = ui.get(f"/cases/{cid}/references", params={"limit": 25}).json()
    selection = ui.get(f"/cases/{cid}/selection").json()["entity_ids"]
    assert ui.get(f"/cases/{cid}/chat").json() == []
    assert ui.get(f"/cases/{cid}/feedback").json() == {"ratings": {}, "flagged": []}
    matching = [r["entity_id"] for r in refs if r["why"].startswith("Matches ")]
    assert selection == matching[:10]  # only matching references are preselected (#110)
    assert [r["why"].startswith("Matches ") for r in refs] == sorted(
        (r["why"].startswith("Matches ") for r in refs), reverse=True)  # matching ones first (#123)
    names = [r["name"] for r in refs]
    assert len(names) == len(set(names)), f"duplicate references: {names}"

    # Review it: the reviewer's score and an internal note
    review = ui.put(f"/cases/{cid}/review", json={"score": 70, "reason": "Strong sensor fit", "author": "IT"}).json()
    assert review["review"]["score"] == 70 and review["score"] == case["score"]  # next to the AI score
    note = ui.post(f"/cases/{cid}/notes", json={"author": "IT", "text": NOTE}).json()

    # Move the card, refine with the chat, keep a chat result, rate one, change the selection
    for status in ["in_review", "accepted"]:
        assert ui.patch(f"/cases/{cid}", json={"status": status}).json()["status"] == status
    chat = ui.post("/chat", json={"session_id": f"case-{cid}", "message": "insulin pump"}).json()
    assert chat["results"]
    new = next((r for r in chat["results"] if r["entity_id"] not in {x["entity_id"] for x in refs}), None)
    if new:
        assert ui.post(f"/cases/{cid}/references", json={"entity_id": new["entity_id"]}).json()["added"] is True
    chosen = [r["entity_id"] for r in refs[:3]] + ([new["entity_id"]] if new else [])
    assert ui.put(f"/cases/{cid}/selection", json={"entity_ids": chosen}).json()["entity_ids"] == chosen
    rating = {"session_id": f"case-{cid}", "query": case["company_name"], "entity_id": refs[0]["entity_id"],
              "rating": "up"}
    assert ui.post("/feedback", json=rating).json() == {"ok": True}

    # Pitch, then report and slides from the selection just saved
    pitch = ui.post(f"/cases/{cid}/pitch", json={"entity_ids": chosen}).json()
    assert pitch["text"] and NOTE not in pitch["text"]
    ui.post(f"/cases/{cid}/documents")
    docs = poll(ui, f"/cases/{cid}/documents", "status")
    assert docs["status"] == "done" and docs["entity_ids"] == chosen
    report = ui.get(f"/cases/{cid}/documents/report")  # the browser's download link
    assert "attachment" in report.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(report.content)).paragraphs)
    assert all(r["name"] in text for r in refs[:3]) and NOTE not in text
    assert len(Presentation(io.BytesIO(ui.get(f"/cases/{cid}/documents/slides").content)).slides) >= 6 + len(chosen)

    # Reload: everything comes back from the backend
    listed = next(c for c in ui.get("/cases").json() if c["id"] == cid)
    assert listed["status"] == "accepted" and listed["review"]["score"] == 70 and listed["notes_count"] == 1
    assert ui.get(f"/cases/{cid}/selection").json()["entity_ids"] == chosen
    assert len(ui.get(f"/cases/{cid}/chat").json()) == 1
    assert ui.get(f"/cases/{cid}/feedback").json()["ratings"] == {str(refs[0]["entity_id"]): "up"}
    assert ui.get(f"/cases/{cid}/pitch").json() == pitch
    assert ui.get(f"/cases/{cid}/documents").json()["status"] == "done"
    assert [n["text"] for n in ui.get(f"/cases/{cid}/notes").json()] == [NOTE]

    # Undo: remove the added reference, the note and the review
    if new:
        assert new["entity_id"] not in ui.delete(f"/cases/{cid}/references/{new['entity_id']}").json()["entity_ids"]
    assert ui.delete(f"/cases/{cid}/notes/{note['id']}").status_code == 200
    assert ui.delete(f"/cases/{cid}/review").json()["review"] is None
    assert ui.get(f"/cases/{cid}/notes").json() == []


def test_sources_and_pages_screens(ui):
    sources = ui.get("/sources").json()
    assert sources
    assert ui.get(f"/sources/{sources[0]['id']}/rules").status_code == 200
    pages = ui.get("/pages").json()
    assert pages and ui.get(f"/pages/{pages[0]['id']}").status_code == 200
