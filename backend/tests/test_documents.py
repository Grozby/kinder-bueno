"""Partner report and slides (#96): checked texts and numbered sources with a fake LLM, the files' content,
and the routes (need Postgres)."""

import asyncio
import io
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from psycopg.types.json import Jsonb

from navigator import documents
from navigator.api import app, get_conn, get_make_documents
from navigator.documents import (
    DocContext,
    SourcePage,
    numbered,
    render_report,
    render_slides,
    write_texts,
)
from navigator.models import Evidence, Result, ScoreItem

NOW = datetime(2026, 10, 18, tzinfo=UTC)


def _ev(claim_id: int, field: str, value: str, url: str, via_unit: str | None = None) -> Evidence:
    return Evidence(claim_id=claim_id, field=field, value=value, quote=f"quote {value}", url=url,
                    judge_verdict=None, judge_reason=None, via_unit=via_unit)


ANNA = Result(
    entity_id=1, type="person", name="Anna Muster", role=["Group Head"], contact=["anna@unibe.example"],
    affiliations=["Diabetes Lab"], expertise=["glucose sensors"], projects=[], infrastructure=[],
    why="Matches 'glucose sensors': Group head who develops glucose sensors [12][13].", score=1.0,
    sources=["https://u/anna"], evidence=[
        _ev(11, "email", "anna@unibe.example", "https://u/anna"), _ev(12, "role", "Group Head", "https://u/anna"),
        _ev(13, "expertise", "glucose sensors", "https://u/anna"),
        _ev(14, "unit_expertise", "type 1 diabetes", "https://u/lab", via_unit="Diabetes Lab"),
    ],
)
LAB = Result(
    entity_id=2, type="unit", name="Diabetes Lab", role=[], contact=[], affiliations=["ARTORG Center"],
    expertise=[], projects=[], infrastructure=[], why="Found for 'clinical studies'. Diabetes research [21].",
    score=0.5, sources=["https://u/lab"], evidence=[_ev(21, "description", "diabetes research", "https://u/lab")],
)


def _ctx(partners=(ANNA, LAB)) -> DocContext:
    urls = ["https://u/anna", "https://u/lab"]
    return DocContext(
        company="Glucora", industry="medical devices", summary="Glucora builds a needle-free glucose patch.",
        needs=["glucose sensors", "clinical studies"], score=72,
        breakdown=[ScoreItem(criterion="glucose sensors", score=90, rationale="Two partners.")],
        partners=list(partners),
        pages={u: SourcePage(u, f"Page {u[-3:]}", "ARTORG Center", NOW) for u in urls},
    )


class FakeLLM:
    """Answers by prompt: the profile prompts list one reference, the summary prompt all of them."""

    def __init__(self, answers: dict[str, object] | None = None, error: Exception | None = None):
        self.answers, self.error = answers or {}, error

    async def chat_json(self, system: str, user: str, max_tokens: int = 4096) -> dict:
        if self.error:
            raise self.error
        kind = "summary" if "Diabetes Lab (unit)" in user and "Anna Muster (person)" in user else \
            "anna" if "Anna Muster (person)" in user else "lab"
        return {"text": self.answers.get(kind)}


# --- texts ---

def test_numbered_turns_claim_ids_into_source_numbers():
    ctx = _ctx()
    assert numbered("Develops glucose sensors [12][13]. In the Diabetes Lab [14].", [ANNA], ctx) == \
        "Develops glucose sensors [1]. In the Diabetes Lab [2]."
    assert numbered("Both [13, 21].", [ANNA, LAB], ctx) == "Both [1, 2]."
    assert numbered("Unknown claim [99].", [ANNA], ctx) == "Unknown claim."


def test_write_texts_keeps_checked_sentences_only():
    llm = FakeLLM({
        "summary": "Glucora needs sensors. Anna Muster develops glucose sensors [13]. Anna Muster builds robots [99].",
        "anna": "Anna Muster is a group head [12]. She develops glucose sensors [13].",
        "lab": "The Diabetes Lab is world famous.",  # names the partner without a citation: nothing left
    })
    summary, profiles = asyncio.run(write_texts(llm, _ctx()))
    assert summary == "Glucora needs sensors. Anna Muster develops glucose sensors [13]."
    assert profiles == {1: "Anna Muster is a group head [12]. She develops glucose sensors [13]."}


def test_write_texts_on_gateway_failure_uses_templates():
    summary, profiles = asyncio.run(write_texts(FakeLLM(error=TimeoutError()), _ctx()))
    assert (summary, profiles) == (None, {})
    assert "Glucora is looking for partners for: glucose sensors; clinical studies." in documents.summary_template(_ctx())
    assert documents.profile_template(ANNA) == \
        "Anna Muster works on: glucose sensors [13]; type 1 diabetes (in Diabetes Lab) [14]."


# --- files ---

def _report_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_report_has_every_section_numbered_citations_and_linked_sources():
    content = render_report(_ctx(), "Anna Muster develops glucose sensors [13].", {}, NOW)
    text = _report_text(content)
    for heading in ["1. Summary", "2. The company", "3. Fit with Canton Bern", "4. Proposed partners",
                    "4.1 Anna Muster", "4.2 Diabetes Lab", "5. Sources"]:
        assert heading in text
    assert "Anna Muster develops glucose sensors [1]." in text
    assert "Why it matches: Matches 'glucose sensors': Group head who develops glucose sensors [1]." in text
    assert "Anna Muster works on: glucose sensors [1]; type 1 diabetes (in Diabetes Lab) [2]." in text  # template
    assert "Score: 72 / 100" in text and "Two partners." in text
    assert "“quote glucose sensors”" in text  # the evidence table
    assert "ARTORG Center: 2 pages, 5 facts about 2 of the partners." in text
    assert "[13]" not in text and "[21]" not in text  # no claim id is left
    links = {r.target_ref for r in Document(io.BytesIO(content)).part.rels.values() if r.is_external}
    assert links == {"https://u/anna", "https://u/lab"}


def test_slides_introduce_the_company_with_one_slide_per_partner():
    prs = Presentation(io.BytesIO(render_slides(_ctx(), NOW)))
    titles = [s.shapes[1].text_frame.text for s in prs.slides]  # shape 0 is the red bar
    assert titles == ["", "About Glucora", "What Glucora is looking for", "Why we are contacting you",
                      "Anna Muster", "Diabetes Lab", "Next steps", "Sources"]
    anna = prs.slides[4]
    text = "\n".join(sh.text_frame.text for sh in anna.shapes if sh.has_text_frame)
    assert "glucose sensors" in text and "type 1 diabetes (in Diabetes Lab)" in text
    assert "Matched need: glucose sensors" in text and "Contact: anna@unibe.example" in text
    assert "[1] https://u/anna" in text
    assert "Group Head" in text and "[12]" not in text
    assert "“quote glucose sensors” (https://u/anna)" in anna.notes_slide.notes_text_frame.text


def test_a_person_with_only_unit_topics_gets_them_under_the_unit_heading():
    only_unit = ANNA.model_copy(update={"evidence": [e for e in ANNA.evidence if e.via_unit]})
    prs = Presentation(io.BytesIO(render_slides(_ctx([only_unit, LAB]), NOW)))
    text = "\n".join(sh.text_frame.text for sh in prs.slides[4].shapes if sh.has_text_frame)
    assert "Research in Diabetes Lab" in text and "•  type 1 diabetes" in text
    assert "Relevant expertise" not in text


def test_files_strip_characters_office_xml_does_not_allow():
    ctx = _ctx()
    ctx.summary = "Glucora\x0b builds\x00 patches."
    assert "Glucora builds patches." in _report_text(render_report(ctx, None, {}, NOW))
    render_slides(ctx, NOW)


def test_sources_continue_on_more_slides():
    ctx = _ctx()
    ctx.pages = {f"https://u/{i}": SourcePage(f"https://u/{i}", None, "DCB", NOW) for i in range(30)}
    prs = Presentation(io.BytesIO(render_slides(ctx, NOW)))
    assert [s.shapes[1].text_frame.text for s in prs.slides][-3:] == \
        ["Sources", "Sources (continued)", "Sources (continued)"]


# --- routes ---

@pytest.fixture
def accepted_case(test_conn, graph) -> str:
    ref = {"entity_id": graph["alice"], "score": 1.0, "need": "glucose sensors", "needs": ["glucose sensors"],
           "why": f"Professor who works on continuous glucose monitoring [{graph['alice_glucose']}]."}
    test_conn.execute(
        """INSERT INTO "case" (id, company_name, file_name, pdf_text, status, analysis_status, summary, needs, refs,
                               selection)
           VALUES ('c1', 'Glucora Health', 'glucora.pdf', 'text', 'accepted', 'done', 'A patch.',
                   '["glucose sensors"]', %s, %s)""",
        (Jsonb([ref]), Jsonb([graph["alice"], graph["lab"]])),  # the lab is selected without being a reference
    )
    return "c1"


@pytest.fixture
def client(test_conn):
    runs = []

    def make(case_id: str) -> None:
        runs.append(case_id)
        documents.generate(test_conn, case_id, write=lambda ctx: (None, {}))  # the gateway is down: templates

    app.dependency_overrides[get_conn] = lambda: test_conn
    app.dependency_overrides[get_make_documents] = lambda: make
    yield TestClient(app), runs
    app.dependency_overrides.clear()


def test_documents_are_generated_stored_and_downloaded(client, graph, accepted_case):
    http, runs = client
    assert http.get(f"/api/cases/{accepted_case}/documents").json()["status"] == "none"
    started = http.post(f"/api/cases/{accepted_case}/documents").json()
    assert started["status"] == "pending" and started["entity_ids"] == [graph["alice"], graph["lab"]]
    assert runs == [accepted_case]  # TestClient runs the background task after the response
    done = http.get(f"/api/cases/{accepted_case}/documents").json()
    assert done["status"] == "done" and done["generated_at"] and done["entity_ids"] == [graph["alice"], graph["lab"]]

    report = http.get(f"/api/cases/{accepted_case}/documents/report")
    assert report.headers["content-type"] == documents.MEDIA_TYPES["report"]
    assert report.headers["content-disposition"] == 'attachment; filename="glucora-health-partner-report.docx"'
    text = _report_text(report.content)
    assert "4.1 Alice Muster" in text and "4.2 Diabetes Lab" in text
    assert "https://artorg.example/alice" in text or "https://artorg.example/alice" in {
        r.target_ref for r in Document(io.BytesIO(report.content)).part.rels.values() if r.is_external}
    assert "surgical robotics" not in text  # judged unsupported: never shown, so never in the report

    slides = http.get(f"/api/cases/{accepted_case}/documents/slides")
    assert slides.headers["content-disposition"] == 'attachment; filename="glucora-health-partner-slides.pptx"'
    assert len(Presentation(io.BytesIO(slides.content)).slides) == 8


def test_documents_only_for_accepted_cases_with_a_selection(client, test_conn, accepted_case):
    http, runs = client
    test_conn.execute("""UPDATE "case" SET status = 'in_review' WHERE id = %s""", (accepted_case,))
    response = http.post(f"/api/cases/{accepted_case}/documents")
    assert (response.status_code, response.json()) == (409, {"detail": "Accept the case first"})
    test_conn.execute("""UPDATE "case" SET status = 'accepted', selection = '[]' WHERE id = %s""", (accepted_case,))
    assert http.post(f"/api/cases/{accepted_case}/documents").status_code == 422
    assert runs == []


def test_a_pending_job_is_returned_not_started_again(client, test_conn, accepted_case):
    http, runs = client
    documents.start(test_conn, accepted_case, [1])
    assert http.post(f"/api/cases/{accepted_case}/documents").json()["status"] == "pending"
    assert runs == []


def test_a_job_pending_for_too_long_shows_as_failed(test_conn, accepted_case):
    old = datetime.now(UTC) - documents.STALE - timedelta(minutes=1)
    test_conn.execute("""UPDATE "case" SET documents = %s WHERE id = %s""",
                      (Jsonb({"status": "pending", "entity_ids": [], "started_at": old.isoformat()}), accepted_case))
    assert documents.status(test_conn, accepted_case).status == "failed"


def test_a_failed_generation_is_stored_as_failed(test_conn, accepted_case):
    documents.start(test_conn, accepted_case, [1])

    def broken(ctx):
        raise RuntimeError("boom")

    documents.generate(test_conn, accepted_case, write=broken)
    found = documents.status(test_conn, accepted_case)
    assert found.status == "failed" and found.error


def test_documents_404(client):
    http, _ = client
    for path in ["/api/cases/nope/documents", "/api/cases/nope/documents/report"]:
        assert http.get(path).status_code == 404
    assert http.post("/api/cases/nope/documents").status_code == 404
    assert http.get("/api/cases/nope/documents/poster").status_code == 422


def test_a_case_deleted_while_its_documents_are_made_stores_nothing(test_conn, accepted_case):
    """#134: the job finds the case gone at the end and stores no files, without failing."""
    documents.start(test_conn, accepted_case, [1])

    def write_then_delete(ctx):
        test_conn.execute('DELETE FROM "case" WHERE id = %s', (accepted_case,))  # the user deleted it meanwhile
        return None, {}

    documents.generate(test_conn, accepted_case, write=write_then_delete)
    assert test_conn.execute("SELECT count(*) FROM case_document").fetchone()[0] == 0
    assert documents.status(test_conn, accepted_case) is None
