"""Partner report (.docx) and slides (.pptx) of an accepted case (#96), from its selected references.

- The report is for the case file and the user: summary, the company, the Canton Bern fit, one section per
  partner and the sources. Apertus writes the summary and one profile per partner, checked like the pitch
  (pitch.check_pitch): a sentence that cites a claim outside the partner's shown claims, or names a partner
  without a citation, is dropped; with nothing left (or on a gateway failure) a template is used.
  Claim ids become numbered references to the Sources chapter, which lists every page with its URL.
- The slides introduce the company to the research partners. Their bullets are stored claim values, never
  LLM text; the quotes and URLs are in the speaker notes.
- Both files are stored as generated: the file that is sent is the file that was checked.
"""

import asyncio
import io
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta

import psycopg
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocInches
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as SlideColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx.util import Pt as SlidePt
from psycopg.types.json import Jsonb

from navigator import analysis, case_state
from navigator.explain import _flat, cited_ids, prompt_claims
from navigator.core.llm import LLM
from navigator.models import CaseDocuments, DocumentKind, Evidence, Result, ScoreItem
from navigator.pitch import check_pitch, format_prompt
from navigator.results import build_results

log = logging.getLogger(__name__)

TIMEOUT = 60.0                  # seconds per text; in the background, so the gateway may take its time
STALE = timedelta(minutes=15)   # a job pending longer was cut off (API restart)
MAX_TOKENS = 900
MAX_REFERENCES = 50             # case references read to find the selected ones
MAX_SLIDE_TOPICS = 5
SOURCES_PER_SLIDE = 12
EVIDENCE_ROWS = 6               # per partner in the report; the most explanatory first (prompt_claims)
REPORT_QUOTE = 220              # characters; the full quote is on the linked page
TOPIC_FIELDS = ("expertise", "unit_expertise", "description")

MEDIA_TYPES = {
    "report": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "slides": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
SUFFIXES = {"report": "partner-report.docx", "slides": "partner-slides.pptx"}

_RULES = """Use ONLY the facts listed for the references. Do not add facts, numbers or names.
Every sentence about a reference MUST cite the claim ids it uses in square brackets, like [12][15].
A sentence about a reference without a claim id is thrown away.
Claim ids are ONLY the numbers in square brackets before each fact, never anything else.
Write natural sentences: do not copy the list, the field names or the quotes.
A "lab topic (via <unit>)" is a topic of the person's unit, NOT the person's own expertise:
write "in <unit>, which works on ..." and name the unit."""

SUMMARY_SYSTEM = f"""You write the summary of a report about a company that looks for research partners in
Canton Bern, for the Bern economic development agency.
Write 2 short paragraphs: what the company needs, then which listed references cover which needs.
{_RULES}
Return ONLY JSON: {{"text": "<paragraphs separated by a blank line>"}}"""

PROFILE_SYSTEM = f"""You write the profile of ONE research partner in Canton Bern for a report about a company.
Write ONE paragraph of 2 to 4 sentences: who the reference is, what it works on, and how that relates to
the company's needs.
{_RULES}
Return ONLY JSON: {{"text": "<paragraph>"}}

Example:
Company: Glucora (medical devices)
Needs: glucose sensors
References:
- Anna Muster (person)
  [12] role: Group Head — "Anna Muster, Group Head"
  [15] expertise: glucose sensors — "develops glucose sensors for type 1 diabetes"
Answer: {{"text": "Anna Muster is a group head [12]. She develops glucose sensors for type 1 diabetes [15], \
which is the sensor know-how Glucora is looking for."}}"""

CITATION_RUN = re.compile(r"\[\d+(?:\s*,\s*\d+)*\](?:\s*\[\d+(?:\s*,\s*\d+)*\])*")
MATCHED_NEED = re.compile(r"^(?:Matches|Found for) '(.+?)'")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # not allowed in Office XML


@dataclass
class SourcePage:
    url: str
    title: str | None
    institution: str
    fetched_at: datetime


@dataclass
class DocContext:
    company: str
    industry: str | None
    summary: str | None
    needs: list[str]
    score: int | None
    breakdown: list[ScoreItem]
    partners: list[Result]
    pages: dict[str, SourcePage]  # by URL, in order of first use; the Sources are numbered in this order

    def number(self, url: str) -> int:
        return list(self.pages).index(url) + 1


# --- Gathering and writing ---

def gather(conn: psycopg.Connection, case_id: str) -> DocContext | None:
    """The case and its selected references (analysis ones and chat-added ones), rebuilt from the current claims."""
    row = conn.execute(
        'SELECT company_name, industry, summary, needs, score, score_breakdown FROM "case" WHERE id = %s', (case_id,)
    ).fetchone()
    if row is None:
        return None
    selection = case_state.get_selection(conn, case_id) or []
    refs = {r.entity_id: r for r in analysis.references(conn, case_id, MAX_REFERENCES) or []}
    other = {r.entity_id: r for r in build_results(conn, [(i, 0.0) for i in selection if i not in refs])}
    partners = [p for i in selection if (p := refs.get(i) or other.get(i))]
    urls = list(dict.fromkeys(e.url for r in partners for e in r.evidence))
    found = {
        url: SourcePage(url, title, institution, fetched)
        for url, title, institution, fetched in conn.execute(
            "SELECT p.url, p.title, s.name, p.fetched_at FROM page p JOIN source s ON s.id = p.source_id "
            "WHERE p.url = ANY(%s)", (urls,),
        )
    }
    return DocContext(
        company=row[0], industry=row[1], summary=row[2], needs=[str(n) for n in row[3]], score=row[4],
        breakdown=[ScoreItem(**b) for b in row[5]], partners=partners,
        pages={u: found[u] for u in urls if u in found},
    )


async def _write(llm: LLM, system: str, prompt: str, results: list[Result], timeout: float) -> str | None:
    try:
        data = await asyncio.wait_for(llm.chat_json(system, prompt, max_tokens=MAX_TOKENS), timeout)
    except Exception as exc:  # noqa: BLE001 - any LLM or network error falls back to the template
        log.warning("document text failed: %s", exc)
        return None
    return check_pitch(data.get("text"), results)


async def write_texts(llm: LLM, ctx: DocContext, timeout: float = TIMEOUT) -> tuple[str | None, dict[int, str]]:
    """The checked summary and the checked profile per partner, all at once; None / missing: use the template."""
    context = (ctx.company, ctx.industry, ctx.summary, ctx.needs)
    texts = await asyncio.gather(
        _write(llm, SUMMARY_SYSTEM, format_prompt(*context, ctx.partners), ctx.partners, timeout),
        *[_write(llm, PROFILE_SYSTEM, format_prompt(*context, [r]), [r], timeout) for r in ctx.partners],
    )
    return texts[0], {r.entity_id: t for r, t in zip(ctx.partners, texts[1:]) if t}


def summary_template(ctx: DocContext) -> str:
    needs = "; ".join(ctx.needs) or "its business model"
    names = ", ".join(r.name for r in ctx.partners)
    return f"{ctx.company} is looking for partners for: {needs}.\n\nThis report proposes {len(ctx.partners)} " \
           f"partners in Canton Bern: {names}."


def _topic(e: Evidence) -> str:
    return f"{_flat(e.value)} (in {e.via_unit})" if e.via_unit else _flat(e.value)


def topics(r: Result) -> list[Evidence]:
    return [e for e in prompt_claims(r.evidence) if e.field in TOPIC_FIELDS]


def profile_template(r: Result) -> str:
    rows = topics(r)[:4]
    if not rows:
        return ""
    return f"{r.name} works on: " + "; ".join(f"{_topic(e)} [{e.claim_id}]" for e in rows) + "."


def numbered(text: str, r: list[Result], ctx: DocContext) -> str:
    """Claim citations ("[12][15]") as numbers of the Sources chapter ("[2, 5]")."""
    urls = {e.claim_id: e.url for x in r for e in x.evidence}

    def sources(m: re.Match) -> str:
        nums = sorted({ctx.number(urls[i]) for i in cited_ids(m.group(0)) if urls.get(i) in ctx.pages})
        return f"[{', '.join(map(str, nums))}]" if nums else ""

    return re.sub(r"\s+([.,;])", r"\1", CITATION_RUN.sub(sources, text)).strip()


def _clean(text: object) -> str:
    return CONTROL.sub("", str(text))


def _filename(company: str, kind: DocumentKind) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", company.lower()).strip("-") or "case"
    return f"{slug}-{SUFFIXES[kind]}"


# --- Report (.docx) ---

def _link(paragraph, url: str) -> None:
    """An external hyperlink run (python-docx has no API for it)."""
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link, run, props = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:rPr")
    link.set(qn("r:id"), rel)
    color, underline = OxmlElement("w:color"), OxmlElement("w:u")
    color.set(qn("w:val"), "0563C1")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = _clean(url)
    text.set(qn("xml:space"), "preserve")
    run.extend([props, text])
    link.append(run)
    paragraph._p.append(link)


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float]):
    """A table with a header row; `widths` in inches, on the columns and on every cell (Word reads the cells,
    other viewers the column grid)."""
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = _clean(text)
    for column, width in zip(table.columns, widths):
        column.width = DocInches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = DocInches(width)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return table


def render_report(ctx: DocContext, summary: str | None, profiles: dict[int, str], now: datetime) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    muted = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading(_clean(f"{ctx.company}: research partners in Canton Bern"), 0)
    meta = doc.add_paragraph()
    meta.add_run(f"Partner report · {now:%d %B %Y} · Research & Expertise Navigator").font.color.rgb = muted
    note = doc.add_paragraph().add_run(
        "Every fact about a partner comes from a public page of its institution and was checked against a quote "
        "from that page. Numbers in brackets, like [3], refer to the Sources in section 5."
    )
    note.italic = True
    note.font.size = Pt(9)

    doc.add_heading("1. Summary", 1)
    for paragraph in (summary or summary_template(ctx)).split("\n\n"):
        doc.add_paragraph(_clean(numbered(paragraph, ctx.partners, ctx)))

    doc.add_heading("2. The company", 1)
    if ctx.industry:
        p = doc.add_paragraph()
        p.add_run("Industry: ").bold = True
        p.add_run(_clean(ctx.industry))
    if ctx.summary:
        doc.add_paragraph(_clean(ctx.summary))
    if ctx.needs:
        doc.add_paragraph().add_run("What the company is looking for").bold = True
        for need in ctx.needs:
            doc.add_paragraph(_clean(need), style="List Bullet")

    doc.add_heading("3. Fit with Canton Bern", 1)
    if ctx.score is not None:
        p = doc.add_paragraph()
        p.add_run("Score: ").bold = True
        p.add_run(f"{ctx.score} / 100, from how well the evidence found covers each need.")
    if ctx.breakdown:
        _table(doc, ["Criterion", "Score", "Rationale"],
               [[b.criterion, str(b.score), b.rationale] for b in ctx.breakdown], [1.6, 0.6, 4.3])

    doc.add_heading("4. Proposed partners", 1)
    for i, r in enumerate(ctx.partners, 1):
        doc.add_heading(_clean(f"4.{i} {r.name}"), 2)
        facts = [("Type", r.type), ("Role", "; ".join(r.role)), ("Part of", "; ".join(r.affiliations)),
                 ("Contact", "; ".join(r.contact))]
        for label, value in facts:
            if value:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.add_run(f"{label}: ").bold = True
                p.add_run(_clean(value))
        why = doc.add_paragraph()
        why.paragraph_format.space_before = Pt(6)
        why.add_run("Why it matches: ").bold = True
        why.add_run(_clean(numbered(r.why, [r], ctx)))
        if text := profiles.get(r.entity_id) or profile_template(r):
            doc.add_paragraph(_clean(numbered(text, [r], ctx)))
        rows = [[f"{e.field}: {_topic(e)}", f"“{_shorten(e.quote, REPORT_QUOTE)}”", f"[{ctx.number(e.url)}]"]
                for e in prompt_claims(r.evidence) if e.url in ctx.pages][:EVIDENCE_ROWS]
        if rows:
            _table(doc, ["Fact", "Quote from the source", "Source"], rows, [2.2, 3.7, 0.6])

    doc.add_heading("5. Sources", 1)
    doc.add_paragraph().add_run("Institutions").bold = True
    for institution in dict.fromkeys(p.institution for p in ctx.pages.values()):
        urls = {u for u, p in ctx.pages.items() if p.institution == institution}
        claims = sum(1 for r in ctx.partners for e in r.evidence if e.url in urls)
        partners = sum(1 for r in ctx.partners if any(e.url in urls for e in r.evidence))
        doc.add_paragraph(
            _clean(f"{institution}: {len(urls)} pages, {claims} facts about {partners} of the partners."),
            style="List Bullet",
        )
    doc.add_paragraph().add_run("Pages").bold = True
    for n, page in enumerate(ctx.pages.values(), 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(_clean(f"[{n}] {page.title or page.institution}. {page.institution}, retrieved "
                         f"{page.fetched_at:%d %B %Y}. "))
        _link(p, page.url)

    doc.add_heading("About this report", 1)
    doc.add_paragraph(
        "The partners were found by the Research & Expertise Navigator on the public pages of the listed "
        "institutions and selected by the reviewer of this case. The summary and the profiles were written by "
        "the Apertus language model; any sentence about a partner that did not cite one of that partner's "
        "checked facts was removed. The company's description comes from the business model it submitted."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Slides (.pptx), for the research partners ---

INK = SlideColor(0x1F, 0x29, 0x37)
ACCENT = SlideColor(0xE3, 0x06, 0x13)  # Bern red, as --bern-red in the UI
GREY = SlideColor(0x6B, 0x72, 0x80)
WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
FONT = "Arial"  # on every computer, so the layout holds


def _text(slide, text: str, left, top, width, height, size: int, color=INK, bold: bool = False):
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = _clean(text)
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = SlidePt(size), bold, color, FONT
    return frame


def _bullets(slide, items: list[str], left, top, width, height, size: int = 18) -> None:
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = SlidePt(8)
        run = p.add_run()
        run.text = _clean(f"•  {item}")
        run.font.size, run.font.color.rgb, run.font.name = SlidePt(size), INK, FONT


def _slide(prs, title: str, notes: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), HEIGHT)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _text(slide, title, Inches(0.8), Inches(0.5), Inches(11.8), Inches(1), 32, bold=True)
    if notes:
        slide.notes_slide.notes_text_frame.text = _clean(notes)
    return slide


def _footer(slide, text: str) -> None:
    _text(slide, text, Inches(0.8), Inches(6.75), Inches(11.8), Inches(0.5), 10, GREY)


def _shorten(text: str, limit: int) -> str:
    text = _flat(text)
    return text if len(text) <= limit else text[: limit - 1].rsplit(" ", 1)[0] + "…"


def render_slides(ctx: DocContext, now: datetime) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    body = (Inches(0.8), Inches(1.7), Inches(11.8), Inches(4.9))

    slide = _slide(prs, "")
    _text(slide, f"Introducing {ctx.company}", Inches(0.8), Inches(2.4), Inches(11.8), Inches(1.2), 44, bold=True)
    _text(slide, "A company looking for research partners in Canton Bern", Inches(0.8), Inches(3.6), Inches(11.8),
          Inches(0.8), 24, GREY)
    _text(slide, " · ".join(filter(None, [ctx.industry, f"{now:%B %Y}"])), Inches(0.8), Inches(4.4), Inches(11.8),
          Inches(0.6), 16, GREY)

    slide = _slide(prs, f"About {ctx.company}")
    if ctx.industry:
        _text(slide, f"Industry: {ctx.industry}", body[0], body[1], body[2], Inches(0.6), 18, GREY)
    _text(slide, _shorten(ctx.summary or "", 900), body[0], Inches(2.4), body[2], Inches(4.2), 20)

    if ctx.needs:
        slide = _slide(prs, f"What {ctx.company} is looking for")
        _bullets(slide, ctx.needs, *body, size=22)

    slide = _slide(prs, "Why we are contacting you")
    _text(slide, f"We matched {ctx.company}'s needs with these research partners in Canton Bern:", body[0], body[1],
          body[2], Inches(0.6), 18, GREY)
    rows = []
    for r in ctx.partners:
        need = MATCHED_NEED.match(r.why)
        unit = r.affiliations[0] if r.affiliations else None
        rows.append(" — ".join(filter(None, [r.name, unit, need and f"for {need.group(1)}"])))
    _bullets(slide, rows[:12], body[0], Inches(2.4), body[2], Inches(4.3), size=16 if len(rows) > 6 else 20)

    for r in ctx.partners:
        rows = topics(r)[:MAX_SLIDE_TOPICS]
        notes = "\n".join(f"{_topic(e)}: “{_flat(e.quote)}” ({e.url})" for e in rows)
        slide = _slide(prs, r.name, notes=f"Facts on this slide and their sources:\n{notes}" if notes else None)
        sub = " · ".join(filter(None, ["; ".join(r.role[:2]), "; ".join(r.affiliations[:2])]))
        if sub:
            _text(slide, sub, body[0], Inches(1.45), body[2], Inches(0.6), 18, GREY)
        # All topics from one unit (#44): say so once, still never as the person's own expertise
        units = {e.via_unit for e in rows}
        unit = next(iter(units)) if len(units) == 1 else None
        heading = f"Research in {unit}" if unit else "Relevant expertise"
        items = [_shorten(_flat(e.value) if unit else _topic(e), 120) for e in rows]
        _text(slide, heading, body[0], Inches(2.2), Inches(7.6), Inches(0.5), 16, ACCENT, bold=True)
        _bullets(slide, items or ["See the sources below."], body[0], Inches(2.7), Inches(7.6), Inches(3.9), size=18)
        side = []
        if need := MATCHED_NEED.match(r.why):
            side.append(f"Matched need: {need.group(1)}")
        if r.contact:
            side.append("Contact: " + "; ".join(r.contact[:2]))
        if side:
            _text(slide, "For " + ctx.company, Inches(8.9), Inches(2.2), Inches(3.7), Inches(0.5), 16, ACCENT,
                  bold=True)
            _bullets(slide, side, Inches(8.9), Inches(2.7), Inches(3.7), Inches(3.9), size=16)
        urls = list(dict.fromkeys(e.url for e in rows)) or list(dict.fromkeys(e.url for e in r.evidence))
        _footer(slide, "Sources: " + "  ".join(f"[{ctx.number(u)}] {u}" for u in urls[:3] if u in ctx.pages))

    slide = _slide(prs, "Next steps")
    _bullets(slide, [
        f"Tell us if a first conversation with {ctx.company} is of interest.",
        "We set up an introduction and share more details with the company's consent.",
        "Questions about the match? Reply to this message.",
    ], *body, size=22)

    pages = list(ctx.pages.values())
    for start in range(0, len(pages), SOURCES_PER_SLIDE):
        slide = _slide(prs, "Sources" if start == 0 else "Sources (continued)")
        chunk = pages[start:start + SOURCES_PER_SLIDE]
        lines = [f"[{start + i}] {p.institution}: {p.url}" for i, p in enumerate(chunk, 1)]
        frame = _text(slide, lines[0], *body, 12)
        for line in lines[1:]:
            run = frame.add_paragraph().add_run()
            run.text = _clean(line)
            run.font.size, run.font.color.rgb, run.font.name = SlidePt(12), INK, FONT
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --- The job ---

def status(conn: psycopg.Connection, case_id: str) -> CaseDocuments | None:
    """The job state; None if no such case. A job pending for too long was cut off by a restart."""
    row = conn.execute('SELECT documents FROM "case" WHERE id = %s', (case_id,)).fetchone()
    if row is None:
        return None
    if row[0] is None:
        return CaseDocuments(status="none")
    docs = CaseDocuments(**row[0])
    if docs.status == "pending" and docs.started_at and datetime.now(UTC) - docs.started_at > STALE:
        return docs.model_copy(update={"status": "failed", "error": "Interrupted (the API restarted). Try again."})
    return docs


def _save(conn: psycopg.Connection, case_id: str, docs: CaseDocuments) -> CaseDocuments:
    conn.execute('UPDATE "case" SET documents = %s WHERE id = %s', (Jsonb(docs.model_dump(mode="json")), case_id))
    return docs


def start(conn: psycopg.Connection, case_id: str, entity_ids: list[int]) -> CaseDocuments:
    return _save(conn, case_id, CaseDocuments(status="pending", entity_ids=entity_ids, started_at=datetime.now(UTC)))


def generate(conn: psycopg.Connection, case_id: str,
             write: Callable[[DocContext], tuple[str | None, dict[int, str]]]) -> None:
    """Make and store both files (the background step after `start`); on any error the job is failed."""
    started = status(conn, case_id)
    try:
        ctx = gather(conn, case_id)
        if ctx is None:
            return
        summary, profiles = write(ctx)
        now = datetime.now(UTC)
        files = {"report": render_report(ctx, summary, profiles, now), "slides": render_slides(ctx, now)}
        with conn.transaction():
            # The case may have been deleted while the texts were written (#134): then store nothing.
            # FOR KEY SHARE keeps a delete from running between this check and the inserts.
            if conn.execute('SELECT 1 FROM "case" WHERE id = %s FOR KEY SHARE', (case_id,)).fetchone() is None:
                log.info("case %s was deleted while its documents were made", case_id)
                return
            for kind, content in files.items():
                conn.execute(
                    "INSERT INTO case_document (case_id, kind, content) VALUES (%s, %s, %s) "
                    "ON CONFLICT (case_id, kind) DO UPDATE SET content = EXCLUDED.content",
                    (case_id, kind, content),
                )
            _save(conn, case_id, CaseDocuments(
                status="done", entity_ids=[r.entity_id for r in ctx.partners],
                started_at=started.started_at if started else None, generated_at=now,
            ))
    except Exception:
        log.exception("case documents failed: %s", case_id)
        _save(conn, case_id, CaseDocuments(
            status="failed", entity_ids=started.entity_ids if started else [],
            started_at=started.started_at if started else None, error="Generating the documents failed. Try again.",
        ))


def load(conn: psycopg.Connection, case_id: str, kind: DocumentKind) -> tuple[bytes, str] | None:
    """The stored file and its download name; None if the case has none (yet)."""
    row = conn.execute(
        'SELECT d.content, c.company_name FROM case_document d JOIN "case" c ON c.id = d.case_id '
        "WHERE d.case_id = %s AND d.kind = %s", (case_id, kind),
    ).fetchone()
    return (bytes(row[0]), _filename(row[1], kind)) if row else None
